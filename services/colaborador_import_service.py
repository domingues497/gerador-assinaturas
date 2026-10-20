"""Importação em lote de colaboradores via XLSX."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook

from models import colaborador as colaborador_model

CAMPOS_OBRIGATORIOS = ["nome", "sobrenome", "cargo"]

MAPEAMENTO_CHAVES: dict[str, list[str]] = {
    "nome": ["nome", "firstname", "first name", "primeiro nome", "primeironome"],
    "sobrenome": ["sobrenome", "lastname", "last name", "ultimo nome", "ultimonome"],
    "cargo": ["cargo", "funcao", "função", "role", "position", "job title", "jobtitle"],
    "departamento": ["departamento", "department", "dept", "depto"],
    "setor": ["setor", "sector", "area", "area"],
    "email": ["email", "e-mail", "e mail", "correio", "mail"],
    "telefone": ["telefone", "tel", "phone", "fixo", "telefone fixo"],
    "ramal": ["ramal", "extension", "ext", "ramal interno"],
    "celular": ["celular", "cel", "mobile", "cellphone", "celular pessoal"],
    "whatsapp": ["whatsapp", "whats", "wa", "zap", "wpp"],
    "endereco": ["endereco", "endereço", "address", "logradouro", "rua"],
    "cidade": ["cidade", "city", "municipio", "município"],
    "estado": ["estado", "state", "uf", "sigla estado"],
    "foto_url": ["foto", "foto_url", "foto url", "imagem", "imagem url", "image", "image url"],
}


def _normalizar_texto(valor: Any) -> str:
    if valor is None:
        return ""
    return str(valor).strip()


def _descobrir_colunas(header_linha: list[Any]) -> dict[str, int]:
    mapa: dict[str, int] = {}
    for idx, celula in enumerate(header_linha):
        chave = _normalizar_texto(celula).lower()
        if not chave:
            continue
        for campo_bd, alias in MAPEAMENTO_CHAVES.items():
            if campo_bd in mapa:
                continue
            if chave in alias:
                mapa[campo_bd] = idx
    return mapa


def _validar_linha(registro: dict[str, str]) -> list[str]:
    erros: list[str] = []
    for campo in CAMPOS_OBRIGATORIOS:
        if not registro.get(campo):
            erros.append(f"Campo obrigatório ausente: {campo}")
    return erros


def importar_xlsx(conteudo_arquivo: bytes) -> dict[str, Any]:
    try:
        wb = load_workbook(io.BytesIO(conteudo_arquivo), data_only=True)
    except Exception as exc:
        return {
            "ok": False,
            "erro": f"Não foi possível ler o arquivo XLSX: {exc}",
            "importados": 0,
            "erros_linha": [],
        }

    if not wb.sheetnames:
        return {
            "ok": False,
            "erro": "A planilha não possui nenhuma aba.",
            "importados": 0,
            "erros_linha": [],
        }

    ws = wb[wb.sheetnames[0]]
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        return {
            "ok": False,
            "erro": "A planilha está vazia.",
            "importados": 0,
            "erros_linha": [],
        }

    cabecalho = linhas[0]
    mapa_colunas = _descobrir_colunas(cabecalho)

    faltantes = [c for c in CAMPOS_OBRIGATORIOS if c not in mapa_colunas]
    if faltantes:
        faltantes_legiveis = ", ".join(faltantes)
        return {
            "ok": False,
            "erro": (
                "Cabeçalho não contém as colunas obrigatórias: "
                f"{faltantes_legiveis}. Utilize cabeçalhos como: Nome, Sobrenome, Cargo."
            ),
            "importados": 0,
            "erros_linha": [],
            "colunas_encontradas": sorted(mapa_colunas.keys()),
        }

    erros_linha: list[dict[str, Any]] = []
    importados: int = 0

    for numero_linha, linha in enumerate(linhas[1:], start=2):
        if all(v is None or _normalizar_texto(v) == "" for v in linha):
            continue
        registro: dict[str, str] = {}
        for campo_bd, idx in mapa_colunas.items():
            if idx < len(linha):
                registro[campo_bd] = _normalizar_texto(linha[idx])
            else:
                registro[campo_bd] = ""
        erros_campo = _validar_linha(registro)
        if erros_campo:
            erros_linha.append({"linha": numero_linha, "erros": erros_campo, "dados": registro})
            continue
        try:
            colaborador_model.criar_colaborador(registro)
            importados += 1
        except Exception as exc:
            erros_linha.append({
                "linha": numero_linha,
                "erros": [f"Erro ao salvar no banco: {exc}"],
                "dados": registro,
            })

    return {
        "ok": importados > 0,
        "erro": None if (importados > 0 or erros_linha) else "Nenhum colaborador válido foi encontrado.",
        "importados": importados,
        "erros_linha": erros_linha,
        "total_processado": len(linhas) - 1,
    }
